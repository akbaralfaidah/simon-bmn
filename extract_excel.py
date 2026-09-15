import sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

# === File 1: Daftar Aset BMN Balai Sumatera ===
wb1 = openpyxl.load_workbook('Daftar Aset BMN Balai Sumatera.xlsx', data_only=True)
ws1 = wb1.active
assets_sumatera = []
for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, values_only=True):
    if not row[0]: continue
    wilayah = str(row[9]) if row[9] else ''
    if 'Sumatera' not in wilayah: continue
    assets_sumatera.append({
        'item_code': str(row[1]) if row[1] else '',
        'name': str(row[2]) if row[2] else '',
        'brand_type': str(row[3]) if row[3] else '',
        'detail': str(row[4]) if row[4] else '',
        'nup': str(int(row[6])) if row[6] else '',
        'destination': str(row[8]) if row[8] else '',
        'value': float(row[10]) if row[10] else 0,
        'vendor': str(row[11]) if row[11] else '',
        'purchase_date': str(row[12]) if row[12] else '',
        'condition': str(row[14]) if row[14] else 'Baik',
        'holder': str(row[17]) if row[17] else '',
    })

# === File 2: Inventaris BMN Jambi ===
wb2 = openpyxl.load_workbook('Inventaris BMN Balai Penegakan Hukum Lingkungan Hidup Jambi.xlsx', data_only=True)
ws2 = wb2.active
assets_jambi = []
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, values_only=True):
    if not row[0] or not row[1]: continue
    assets_jambi.append({
        'name': str(row[1]) if row[1] else '',
        'coordinates': str(row[5]) if row[5] else '',
        'condition': str(row[6]) if row[6] else 'Baik',
    })

# === File 3: Laptop ===
wb3 = openpyxl.load_workbook('DAFTAR BMN LAPTOP.xlsx', data_only=True)
ws3 = wb3.active
laptops = []
for row in ws3.iter_rows(min_row=6, max_row=ws3.max_row, values_only=True):
    if not row[0] or not row[1]: continue
    nama = str(row[1]) if row[1] else ''
    nip = str(row[2]) if row[2] else ''
    jabatan = str(row[4]) if row[4] else ''
    laptop_lama = str(row[6]) if row[6] else ''
    nup_lama = str(row[7]) if row[7] else ''
    laptop_baru = str(row[9]) if row[9] else ''
    nup_baru = str(row[10]) if row[10] else ''
    if laptop_lama or laptop_baru:
        laptops.append({
            'holder_name': nama,
            'holder_nip': nip,
            'holder_jabatan': jabatan,
            'laptop_old': laptop_lama,
            'nup_old': nup_lama,
            'laptop_new': laptop_baru,
            'nup_new': nup_baru,
        })

print(f"Sumatera assets: {len(assets_sumatera)}")
print(f"Jambi assets: {len(assets_jambi)}")
print(f"Laptops: {len(laptops)}")

# Save as JSON for PHP seeder
with open('database/data/bmn_sumatera.json', 'w', encoding='utf-8') as f:
    json.dump(assets_sumatera, f, ensure_ascii=False, indent=2, default=str)
with open('database/data/bmn_jambi.json', 'w', encoding='utf-8') as f:
    json.dump(assets_jambi, f, ensure_ascii=False, indent=2, default=str)
with open('database/data/bmn_laptops.json', 'w', encoding='utf-8') as f:
    json.dump(laptops, f, ensure_ascii=False, indent=2, default=str)

print("JSON files saved to database/data/")
